import re
import shutil
from pathlib import Path

import pytest
from docx.shared import Inches
from docxtpl import DocxTemplate

from pl2docx.element_config import ElementConfig, FillInPreferences, SelectorPreferences
from pl2docx.element_renderer import build_question_context, create_list_formats
from pl2docx.html_parser import ImageRef, MathRef, ParsedQuestion, SvgRef, Widget, plain
from pl2docx.latex_math import LatexRenderError
from pl2docx.svg_render import SvgRenderError

from conftest import chromium_available

FIXTURES_DIR = Path(__file__).parent / "fixtures" / "instance_question"


def _tpl(starter_template) -> DocxTemplate:
    tpl = DocxTemplate(str(starter_template))
    tpl.init_docx()
    return tpl


def _mc_question(prompt=("Pick one.", "")):
    widget = Widget(
        kind="multiple_choice",
        name="statement",
        options=[plain("Alpha"), plain("Beta")],
        correct_option_indices=[1],
        is_inline=True,
    )
    return ParsedQuestion(
        title="Q",
        prompt_segments=[plain(p) for p in prompt],
        widgets=[widget],
        answer_panel_text=plain("(B) Beta"),
        points="1",
        points_numeric=1.0,
        qid="q/1",
    )


def _string_question(label=None, suffix=None):
    widget = Widget(
        kind="string_input",
        name="answer",
        label=plain(label) if label is not None else None,
        suffix=plain(suffix) if suffix is not None else None,
    )
    return ParsedQuestion(
        title="Q",
        prompt_segments=[plain("Fill in:"), plain("")],
        widgets=[widget],
        answer_panel_text=None,
        points="1",
        points_numeric=1.0,
        qid="q/2",
    )


def _text(subdoc) -> str:
    return "\n".join(p.text for p in subdoc.paragraphs)


def _has_run_border(subdoc) -> bool:
    return any("w:bdr" in run._element.xml for p in subdoc.paragraphs for run in p.runs)


def _num_id_of(paragraph) -> int | None:
    match = re.search(r'<w:numId w:val="(\d+)"/>', paragraph._p.xml)
    return int(match.group(1)) if match else None


def _num_fmt_for(tpl: DocxTemplate, num_id: int) -> str:
    """Look up the real numFmt ("decimal"/"bullet") a paragraph's numId resolves to.

    Parses numbering.xml's raw XML string rather than using `.xpath()`
    directly - python-docx's `BaseOxmlElement.xpath()` auto-injects its own
    fixed namespace map and rejects an explicit `namespaces=` kwarg, but a
    `<w:abstractNum>` element (no python-docx object-model support at all -
    see element_renderer.py) comes back as a plain, differently-behaved lxml
    element instead once selected via xpath, needing its own explicit
    `namespaces=`. Regexing the serialized XML sidesteps that inconsistency.
    """
    xml = tpl.get_docx().part.numbering_part.element.xml
    num_match = re.search(rf'<w:num w:numId="{num_id}">\s*<w:abstractNumId w:val="(\d+)"', xml)
    assert num_match is not None
    abstract_id = num_match.group(1)
    abstract_match = re.search(
        rf'<w:abstractNum w:abstractNumId="{abstract_id}">.*?<w:numFmt w:val="(\w+)"', xml, re.DOTALL
    )
    assert abstract_match is not None
    return abstract_match.group(1)


def test_selector_inline_display_bold_correct(starter_template):
    tpl = _tpl(starter_template)
    element_config = ElementConfig(
        preferences={"multiple_choice": SelectorPreferences(list_style="letter-labels", display="inline")},
        behavior_class={},
    )
    ctx = build_question_context(tpl, _mc_question(), 1, element_config)
    text = _text(ctx["question_contents"])
    assert "(A) Alpha" in text
    assert "(B) Beta" in text
    assert "Pick one." in text
    # all on one paragraph (inline)
    assert len(ctx["question_contents"].paragraphs) == 1

    beta_run = next(r for p in ctx["question_contents"].paragraphs for r in p.runs if "Beta" in r.text)
    assert beta_run.bold is True


def test_selector_block_display_one_option_per_line(starter_template):
    tpl = _tpl(starter_template)
    element_config = ElementConfig(
        preferences={"multiple_choice": SelectorPreferences(list_style="bubble", display="block")},
        behavior_class={},
    )
    ctx = build_question_context(tpl, _mc_question(), 1, element_config)
    paragraphs = [p.text for p in ctx["question_contents"].paragraphs]
    assert any(p.strip() == "○ Alpha" for p in paragraphs)
    assert any(p.strip() == "○ Beta" for p in paragraphs)


def test_selector_draw_border_inline_uses_run_border(starter_template):
    tpl = _tpl(starter_template)
    element_config = ElementConfig(
        preferences={"multiple_choice": SelectorPreferences(display="inline", draw_border=True)},
        behavior_class={},
    )
    ctx = build_question_context(tpl, _mc_question(), 1, element_config)
    assert _has_run_border(ctx["question_contents"])
    # no table should be created for inline draw-border
    assert len(ctx["question_contents"].tables) == 0


def test_selector_draw_border_block_uses_run_border_not_table(starter_template):
    """draw-border must stay a tight, text-width run border for every display mode
    (including "block", one option per line) — a table-cell border was tried and
    rejected: it always spans the full page width regardless of the boxed text's
    actual width."""
    tpl = _tpl(starter_template)
    element_config = ElementConfig(
        preferences={"multiple_choice": SelectorPreferences(display="block", draw_border=True)},
        behavior_class={},
    )
    ctx = build_question_context(tpl, _mc_question(), 1, element_config)
    assert len(ctx["question_contents"].tables) == 0
    assert _has_run_border(ctx["question_contents"])
    paragraphs = [p.text for p in ctx["question_contents"].paragraphs]
    assert any("Alpha" in p for p in paragraphs)
    assert any("Beta" in p for p in paragraphs)


def test_fill_in_no_label_falls_back_to_answer_prefix(starter_template):
    tpl = _tpl(starter_template)
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, _string_question(), 1, element_config)
    text = _text(ctx["question_contents"])
    assert "Answer: " in text


def test_fill_in_label_and_suffix_used(starter_template):
    tpl = _tpl(starter_template)
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, _string_question(label="pH =", suffix="units"), 1, element_config)
    text = _text(ctx["question_contents"])
    assert "pH =" in text
    assert "units" in text
    assert "Answer:" not in text


def test_fill_in_template_display_routes_to_answer_element(starter_template):
    tpl = _tpl(starter_template)
    element_config = ElementConfig(
        preferences={"string_input": FillInPreferences(display="template")},
        behavior_class={},
    )
    ctx = build_question_context(tpl, _string_question(label="pH =", suffix="units"), 1, element_config)
    assert "pH =" not in _text(ctx["question_contents"])
    assert "pH =" in _text(ctx["answer_element"])


def test_fill_in_none_display_generated_then_discarded(starter_template):
    tpl = _tpl(starter_template)
    element_config = ElementConfig(
        preferences={"string_input": FillInPreferences(display="none")},
        behavior_class={},
    )
    ctx = build_question_context(tpl, _string_question(label="pH ="), 1, element_config)
    assert "pH =" not in _text(ctx["question_contents"])
    assert _text(ctx["answer_element"]).strip() == ""


def test_additional_element_fill_in_renders_end_to_end(starter_template):
    """A widget of a non-built-in kind (as produced for an additional-elements
    tag, e.g. pl-scinum-input) must render exactly like any other fill-in
    widget - build_question_context/element_renderer are kind-agnostic beyond
    the built-in selector-kind check."""
    tpl = _tpl(starter_template)
    widget = Widget(kind="pl-scinum-input", name="first", label=plain("98.0:"))
    question = ParsedQuestion(
        title="Q",
        prompt_segments=[plain("Part A:"), plain("")],
        widgets=[widget],
        answer_panel_text=None,
        points="1",
        points_numeric=1.0,
        qid="q/3",
    )
    element_config = ElementConfig(
        preferences={"pl-scinum-input": FillInPreferences(draw_border=True)},
        behavior_class={"pl-scinum-input": "fill-in"},
    )
    ctx = build_question_context(tpl, question, 1, element_config)
    text = _text(ctx["question_contents"])
    assert "98.0:" in text
    assert "Answer:" not in text
    assert _has_run_border(ctx["question_contents"])


def test_widgets_render_at_source_position():
    """Regression: widget content must appear between its own prompt segments,
    not be appended after the whole prompt."""
    from pl2docx.html_parser import parse_instance_question_html, plain_text

    html = """
    <div class="question-block">
      <div class="card-header"><h1>Ordering</h1></div>
      <div class="card-body question-body">
        <p>Text A</p>
        <span class="input-group pl-string-input">
          <input class="form-control pl-string-input-input" name="answer" type="text">
        </span>
        <p>Text B</p>
      </div>
    </div>
    """
    q = parse_instance_question_html(html)
    assert [plain_text(seg) for seg in q.prompt_segments] == ["Text A", "Text B"]


def test_rich_text_fixture_renders_end_to_end(starter_template):
    """The rich_text_key.html fixture's bold/italic answer-panel content must
    survive all the way through build_question_context into real docx runs."""
    from pathlib import Path

    from pl2docx.html_parser import parse_instance_question_html

    fixtures_dir = Path(__file__).parent / "fixtures" / "instance_question"
    html = (fixtures_dir / "rich_text_key.html").read_text(encoding="utf-8")
    question = parse_instance_question_html(html)

    tpl = _tpl(starter_template)
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, question, 1, element_config)

    answer_text = _text(ctx["answer_contents"])
    assert "limiting reagent" in answer_text
    assert "Oxygen: 0 mol remaining" in answer_text

    bold_run = next(
        r
        for p in ctx["answer_contents"].paragraphs
        for r in p.runs
        if "limiting reagent" in r.text
    )
    assert bold_run.bold is True
    italic_run = next(
        r for p in ctx["answer_contents"].paragraphs for r in p.runs if r.text.strip() == "oxygen"
    )
    assert italic_run.italic is True


def test_prompt_bold_run_rendered_with_real_bold_formatting(starter_template):
    """Phase 4 increment 1: a bold TextRun node in a prompt segment must produce
    a real bold run in the rendered subdoc, not just plain concatenated text."""
    from pl2docx.html_parser import TextRun

    tpl = _tpl(starter_template)
    widget = Widget(kind="integer_input", name="answer")
    question = ParsedQuestion(
        title="Q",
        prompt_segments=[
            [TextRun("This is "), TextRun("bold", bold=True), TextRun(" text.")],
            plain(""),
        ],
        widgets=[widget],
        answer_panel_text=None,
        points="1",
        points_numeric=1.0,
        qid="q/4",
    )
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, question, 1, element_config)
    text = _text(ctx["question_contents"])
    assert "This is bold text." in text

    bold_run = next(
        r for p in ctx["question_contents"].paragraphs for r in p.runs if r.text.strip() == "bold"
    )
    assert bold_run.bold is True
    non_bold_runs = [
        r for p in ctx["question_contents"].paragraphs for r in p.runs if "This is" in r.text
    ]
    assert all(not r.bold for r in non_bold_runs)


def _inline_shape_count(subdoc) -> int:
    """Count real embedded pictures in `subdoc`.

    `Subdoc.inline_shapes` (a docxtpl passthrough to the underlying
    python-docx `Document`) doesn't reliably reflect pictures added via
    `Run.add_picture()` on paragraphs docxtpl created - confirmed empty even
    when a `<w:drawing>` element genuinely is present in a run's XML.
    Checking the run XML directly is what actually works.
    """
    return sum(
        1
        for p in subdoc.paragraphs
        for r in p.runs
        if "<w:drawing>" in r._element.xml
    )


def test_image_embeds_as_real_picture_when_base_dir_given(starter_template):
    """Phase 4 increment 2: an ImageRef must become a real embedded picture
    (not the "[image]" fallback text) when image_base_dir resolves to a real
    file - uses the image_blank.html fixture + its sibling files/tiny_dot.png."""
    from pl2docx.html_parser import parse_instance_question_html

    html = (FIXTURES_DIR / "image_blank.html").read_text(encoding="utf-8")
    question = parse_instance_question_html(html)

    tpl = _tpl(starter_template)
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, question, 1, element_config, image_base_dir=FIXTURES_DIR)

    assert _inline_shape_count(ctx["question_contents"]) == 1
    assert "[image]" not in _text(ctx["question_contents"])

    # width="150" in the source <img> -> 150/96 inch, in EMU - checked via the
    # raw <wp:extent cx="..."> XML since Subdoc.inline_shapes isn't reliable
    # here (see _inline_shape_count).
    picture_run = next(
        r
        for p in ctx["question_contents"].paragraphs
        for r in p.runs
        if "<w:drawing>" in r._element.xml
    )
    match = re.search(r'<wp:extent cx="(\d+)"', picture_run._element.xml)
    assert match is not None
    assert int(match.group(1)) == int(Inches(150 / 96))


def test_image_falls_back_to_alt_text_without_image_base_dir(starter_template):
    """No image_base_dir given (the pre-increment-2 default) must keep the
    old alt-text fallback behavior, not attempt to embed anything."""
    from pl2docx.html_parser import parse_instance_question_html

    html = (FIXTURES_DIR / "image_blank.html").read_text(encoding="utf-8")
    question = parse_instance_question_html(html)

    tpl = _tpl(starter_template)
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, question, 1, element_config)

    assert _inline_shape_count(ctx["question_contents"]) == 0
    assert "a red dot" in _text(ctx["question_contents"])


def test_image_falls_back_to_alt_text_when_file_missing(starter_template):
    """A referenced image that isn't actually on disk (e.g. an external image
    fetch.py deliberately didn't download) must fall back gracefully, not
    raise and fail the whole render."""
    tpl = _tpl(starter_template)
    widget = Widget(kind="integer_input", name="answer")
    question = ParsedQuestion(
        title="Q",
        prompt_segments=[
            [ImageRef(local_path="files/does_not_exist.png", alt="missing image")],
            plain(""),
        ],
        widgets=[widget],
        answer_panel_text=None,
        points="1",
        points_numeric=1.0,
        qid="q/5",
    )
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, question, 1, element_config, image_base_dir=FIXTURES_DIR)

    assert _inline_shape_count(ctx["question_contents"]) == 0
    assert "missing image" in _text(ctx["question_contents"])


def _math_question(latex="x^2", display_mode=False):
    widget = Widget(kind="integer_input", name="answer")
    return ParsedQuestion(
        title="Q",
        prompt_segments=[
            [MathRef(latex=latex, display_mode=display_mode)],
            plain(""),
        ],
        widgets=[widget],
        answer_panel_text=None,
        points="1",
        points_numeric=1.0,
        qid="q/6",
    )


@pytest.mark.skipif(
    shutil.which("latex") is None or shutil.which("dvipng") is None,
    reason="requires a LaTeX install (latex) and dvipng on PATH",
)
def test_math_ref_embeds_as_real_picture(starter_template):
    """Phase 4 increment 3: a MathRef must become a real embedded (LaTeX-
    rendered) picture, not the placeholder "$latex$" text, when a LaTeX
    install is available."""
    tpl = _tpl(starter_template)
    question = _math_question("x^2 + y^2 = z^2")
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, question, 1, element_config)

    assert _inline_shape_count(ctx["question_contents"]) == 1
    assert "$x^2 + y^2 = z^2$" not in _text(ctx["question_contents"])


def test_math_ref_falls_back_to_placeholder_text_on_render_error(starter_template, monkeypatch):
    """A LaTeX compile failure (or no LaTeX install at all) must degrade to
    the old raw-text placeholder, not crash the whole document render."""
    import pl2docx.element_renderer as element_renderer_module

    def _boom(latex, display_mode, **kwargs):
        raise LatexRenderError("simulated failure")

    monkeypatch.setattr(element_renderer_module, "render_math_png", _boom)

    tpl = _tpl(starter_template)
    question = _math_question("x^2", display_mode=False)
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, question, 1, element_config)

    assert _inline_shape_count(ctx["question_contents"]) == 0
    assert "$x^2$" in _text(ctx["question_contents"])


def _mock_render_math_png(monkeypatch):
    """Stub out real LaTeX compilation with the existing tiny_dot.png fixture,
    so display-math paragraph-isolation tests don't need a LaTeX install."""
    import pl2docx.element_renderer as element_renderer_module
    from pl2docx.latex_math import RenderedMath

    fake = RenderedMath(png_path=FIXTURES_DIR / "files" / "tiny_dot.png", width_in=0.5, depth_pt=0.0)
    monkeypatch.setattr(element_renderer_module, "render_math_png", lambda *a, **kw: fake)


def _break_count(paragraph) -> int:
    return sum(r._element.xml.count("<w:br/>") for r in paragraph.runs)


def test_display_math_separated_by_soft_breaks(starter_template, monkeypatch):
    """Display-mode math ($$...$$) must be separated from surrounding text by
    soft line breaks (Run.add_break()) within the SAME paragraph, not
    isolated into its own real paragraph - the user chose this specifically
    so the same mechanism works unchanged for display math inside a list
    item (real content), rather than needing separate list/non-list logic."""
    _mock_render_math_png(monkeypatch)
    tpl = _tpl(starter_template)
    widget = Widget(kind="integer_input", name="answer")
    question = ParsedQuestion(
        title="Q",
        prompt_segments=[
            [
                *plain("Before."),
                MathRef(latex="a+b", display_mode=True),
                *plain("After."),
            ],
            plain(""),
        ],
        widgets=[widget],
        answer_panel_text=None,
        points="1",
        points_numeric=1.0,
        qid="q/7",
    )
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, question, 1, element_config)

    paragraphs = [p for p in ctx["question_contents"].paragraphs if p.text.strip() or p.runs]
    math_paragraphs = [
        p for p in paragraphs if any("<w:drawing>" in r._element.xml for r in p.runs)
    ]
    math_paragraph = math_paragraphs[0]
    # "Before."/math/"After." must all share the same paragraph - not split
    # across separate ones - even though this question also has its own
    # (unrelated) widget paragraph elsewhere in question_contents.
    assert "Before." in math_paragraph.text
    assert "After." in math_paragraph.text
    assert _break_count(math_paragraph) == 2  # one before, one after the image


def test_display_math_no_break_when_alone(starter_template, monkeypatch):
    """No soft break at all when display math is already the only (first and
    last) thing in its segment."""
    _mock_render_math_png(monkeypatch)
    tpl = _tpl(starter_template)
    widget = Widget(kind="integer_input", name="answer")
    question = ParsedQuestion(
        title="Q",
        prompt_segments=[
            [MathRef(latex="a+b", display_mode=True)],
            plain(""),
        ],
        widgets=[widget],
        answer_panel_text=None,
        points="1",
        points_numeric=1.0,
        qid="q/8",
    )
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, question, 1, element_config)

    paragraphs = [p for p in ctx["question_contents"].paragraphs if p.text.strip() or p.runs]
    math_paragraphs = [
        p for p in paragraphs if any("<w:drawing>" in r._element.xml for r in p.runs)
    ]
    assert len(math_paragraphs) == 1
    assert _break_count(math_paragraphs[0]) == 0


def test_list_item_marker_shares_line_with_widget_content(starter_template):
    """Regression: a widget that's the sole content of a <li> (real content:
    every pl-scinum-input item in TEST/pl-scinum-input) must render with its
    "1. "/"2. " marker on the SAME line as the widget, not orphaned alone on
    its own paragraph above an unmarked widget line."""
    from pl2docx.html_parser import parse_instance_question_html

    html = """
    <div class="question-block">
      <div class="card-header"><h1>List of widgets</h1></div>
      <div class="card-body question-body">
        <ol>
          <li>
            <span class="input-group pl-scinum-input">
              <span class="input-group-text">first:</span>
              <input class="form-control pl-scinum-input-input" name="a" type="text">
            </span>
          </li>
          <li>
            <span class="input-group pl-scinum-input">
              <span class="input-group-text">second:</span>
              <input class="form-control pl-scinum-input-input" name="b" type="text">
            </span>
          </li>
        </ol>
      </div>
    </div>
    """
    question = parse_instance_question_html(html, additional_fill_in_tags=["pl-scinum-input"])
    tpl = _tpl(starter_template)
    element_config = ElementConfig(preferences={}, behavior_class={"pl-scinum-input": "fill-in"})
    ctx = build_question_context(tpl, question, 1, element_config)

    paragraphs = [p.text.strip() for p in ctx["question_contents"].paragraphs if p.text.strip()]
    assert any(p.startswith("1.") and "first:" in p for p in paragraphs)
    assert any(p.startswith("2.") and "second:" in p for p in paragraphs)
    # No orphaned "1."/"2."-only paragraph separate from its widget content.
    assert not any(p in ("1.", "2.") for p in paragraphs)


def test_list_item_marker_survives_with_inline_selector_widget(starter_template):
    """Same regression, but for an inline-display selector widget (real
    content: question 235's numbered image-pair-plus-classification-choice
    items) rather than a block-display fill-in widget."""
    from pl2docx.html_parser import parse_instance_question_html

    html = """
    <div class="question-block">
      <div class="card-header"><h1>Numbered choices</h1></div>
      <div class="card-body question-body">
        <ol>
          <li>Classify item one:
            <div class="form-check form-check-inline">
              <input type="radio" name="q1" value="a">
              <label><div class="pl-multiple-choice-answer">Physical</div></label>
            </div>
            <div class="form-check form-check-inline">
              <input type="radio" name="q1" value="b">
              <label><div class="pl-multiple-choice-answer">Chemical</div></label>
            </div>
          </li>
          <li>Classify item two:
            <div class="form-check form-check-inline">
              <input type="radio" name="q2" value="a">
              <label><div class="pl-multiple-choice-answer">Physical</div></label>
            </div>
            <div class="form-check form-check-inline">
              <input type="radio" name="q2" value="b">
              <label><div class="pl-multiple-choice-answer">Chemical</div></label>
            </div>
          </li>
        </ol>
      </div>
    </div>
    """
    question = parse_instance_question_html(html)
    tpl = _tpl(starter_template)
    element_config = ElementConfig(
        preferences={"multiple_choice": SelectorPreferences(display="inline")},
        behavior_class={},
    )
    ctx = build_question_context(tpl, question, 1, element_config)

    paragraphs = [p.text.strip() for p in ctx["question_contents"].paragraphs if p.text.strip()]
    assert any(p.startswith("1.") and "Classify item one" in p and "Physical" in p for p in paragraphs)
    assert any(p.startswith("2.") and "Classify item two" in p and "Physical" in p for p in paragraphs)


def test_real_ordered_list_uses_numpr_with_decimal_format(starter_template):
    """Phase 4 increment 1 follow-up: with list_formats supplied, list items
    must be real Word list paragraphs (<w:numPr>, resolving to a "decimal"
    numFmt) - not plain prepended "1. " text."""
    from pl2docx.html_parser import parse_instance_question_html

    html = """
    <div class="question-block">
      <div class="card-header"><h1>Ordered list</h1></div>
      <div class="card-body question-body">
        <ol>
          <li>First item</li>
          <li>Second item</li>
        </ol>
        <span class="input-group pl-integer-input">
          <input class="form-control pl-integer-input-input" name="answer" type="text">
        </span>
      </div>
    </div>
    """
    question = parse_instance_question_html(html)
    tpl = _tpl(starter_template)
    list_formats = create_list_formats(tpl)
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, question, 1, element_config, list_formats=list_formats)

    numbered = [p for p in ctx["question_contents"].paragraphs if _num_id_of(p) is not None]
    assert len(numbered) == 2
    assert "First item" not in [p.text for p in ctx["question_contents"].paragraphs if _num_id_of(p) is None]
    for p in numbered:
        assert "1." not in p.text and "2." not in p.text  # no more fake prepended text
        assert _num_fmt_for(tpl, _num_id_of(p)) == "decimal"
    # Both items belong to the same list -> same numId, so Word counts them 1, 2.
    assert _num_id_of(numbered[0]) == _num_id_of(numbered[1])


def test_real_unordered_list_uses_bullet_format(starter_template):
    from pl2docx.html_parser import parse_instance_question_html

    html = """
    <div class="question-block">
      <div class="card-header"><h1>Unordered list</h1></div>
      <div class="card-body question-body">
        <ul>
          <li>Alpha</li>
          <li>Beta</li>
        </ul>
        <span class="input-group pl-integer-input">
          <input class="form-control pl-integer-input-input" name="answer" type="text">
        </span>
      </div>
    </div>
    """
    question = parse_instance_question_html(html)
    tpl = _tpl(starter_template)
    list_formats = create_list_formats(tpl)
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, question, 1, element_config, list_formats=list_formats)

    numbered = [p for p in ctx["question_contents"].paragraphs if _num_id_of(p) is not None]
    assert len(numbered) == 2
    for p in numbered:
        assert _num_fmt_for(tpl, _num_id_of(p)) == "bullet"


def test_multiline_list_item_stays_one_paragraph_with_soft_break(starter_template):
    """Regression (the bug reported from real content): a <li> with multiple
    lines (here: two <p>s) must render as ONE Word list paragraph with a soft
    line break between them, not two separate paragraphs where only the
    first is indented/numbered."""
    from pl2docx.html_parser import parse_instance_question_html

    html = """
    <div class="question-block">
      <div class="card-header"><h1>Multi-line item</h1></div>
      <div class="card-body question-body">
        <ol>
          <li><p>First line of item one.</p><p>Second line of item one.</p></li>
          <li>Item two.</li>
        </ol>
        <span class="input-group pl-integer-input">
          <input class="form-control pl-integer-input-input" name="answer" type="text">
        </span>
      </div>
    </div>
    """
    question = parse_instance_question_html(html)
    tpl = _tpl(starter_template)
    list_formats = create_list_formats(tpl)
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, question, 1, element_config, list_formats=list_formats)

    numbered = [p for p in ctx["question_contents"].paragraphs if _num_id_of(p) is not None]
    assert len(numbered) == 2  # one real numbered paragraph per <li>, not per line
    item_one = numbered[0]
    assert "First line of item one." in item_one.text
    assert "Second line of item one." in item_one.text
    assert "<w:br/>" in item_one._p.xml  # soft break between the two lines


def test_list_split_across_widget_interleaving_keeps_same_numid(starter_template):
    """Real content (physical-or-chemical): a single <ol> whose items each
    contain a widget ends up split across multiple prompt_segments by the
    widget markers - all items must still share one numId (same list,
    continuous 1/2/3 numbering), not restart or diverge per segment."""
    from pl2docx.html_parser import parse_instance_question_html

    html = """
    <div class="question-block">
      <div class="card-header"><h1>Split list</h1></div>
      <div class="card-body question-body">
        <ol>
          <li>Item one:
            <span class="input-group pl-integer-input">
              <input class="form-control pl-integer-input-input" name="a" type="text">
            </span>
          </li>
          <li>Item two:
            <span class="input-group pl-integer-input">
              <input class="form-control pl-integer-input-input" name="b" type="text">
            </span>
          </li>
        </ol>
      </div>
    </div>
    """
    question = parse_instance_question_html(html)
    tpl = _tpl(starter_template)
    list_formats = create_list_formats(tpl)
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, question, 1, element_config, list_formats=list_formats)

    numbered = [p for p in ctx["question_contents"].paragraphs if _num_id_of(p) is not None]
    assert len(numbered) == 2
    assert _num_id_of(numbered[0]) == _num_id_of(numbered[1])
    assert "Item one:" in numbered[0].text
    assert "Item two:" in numbered[1].text


def test_independent_lists_get_different_num_ids(starter_template):
    """Two separate <ol>s in the same question must each restart their own
    numbering (different Word numId instances), not share one continuous
    count."""
    from pl2docx.html_parser import parse_instance_question_html

    html = """
    <div class="question-block">
      <div class="card-header"><h1>Two lists</h1></div>
      <div class="card-body question-body">
        <ol><li>List A item one</li><li>List A item two</li></ol>
        <ol><li>List B item one</li><li>List B item two</li></ol>
        <span class="input-group pl-integer-input">
          <input class="form-control pl-integer-input-input" name="answer" type="text">
        </span>
      </div>
    </div>
    """
    question = parse_instance_question_html(html)
    tpl = _tpl(starter_template)
    list_formats = create_list_formats(tpl)
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, question, 1, element_config, list_formats=list_formats)

    numbered = [p for p in ctx["question_contents"].paragraphs if _num_id_of(p) is not None]
    assert len(numbered) == 4
    list_a_ids = {_num_id_of(numbered[0]), _num_id_of(numbered[1])}
    list_b_ids = {_num_id_of(numbered[2]), _num_id_of(numbered[3])}
    assert list_a_ids == {_num_id_of(numbered[0])}  # both List A items share a numId
    assert list_b_ids == {_num_id_of(numbered[2])}  # both List B items share a numId
    assert list_a_ids != list_b_ids  # but the two lists don't share one


def test_widget_sole_content_of_li_gets_real_list_numbering(starter_template):
    """Real content (TEST/pl-scinum-input): a <li> whose sole content is a
    widget must have real <w:numPr> applied to the paragraph holding that
    widget's content, not just plain "1. " text prepended somewhere."""
    from pl2docx.html_parser import parse_instance_question_html

    html = """
    <div class="question-block">
      <div class="card-header"><h1>List of widgets</h1></div>
      <div class="card-body question-body">
        <ol>
          <li>
            <span class="input-group pl-scinum-input">
              <span class="input-group-text">first:</span>
              <input class="form-control pl-scinum-input-input" name="a" type="text">
            </span>
          </li>
          <li>
            <span class="input-group pl-scinum-input">
              <span class="input-group-text">second:</span>
              <input class="form-control pl-scinum-input-input" name="b" type="text">
            </span>
          </li>
        </ol>
      </div>
    </div>
    """
    question = parse_instance_question_html(html, additional_fill_in_tags=["pl-scinum-input"])
    tpl = _tpl(starter_template)
    list_formats = create_list_formats(tpl)
    element_config = ElementConfig(preferences={}, behavior_class={"pl-scinum-input": "fill-in"})
    ctx = build_question_context(tpl, question, 1, element_config, list_formats=list_formats)

    numbered = [p for p in ctx["question_contents"].paragraphs if _num_id_of(p) is not None]
    assert len(numbered) == 2
    assert "first:" in numbered[0].text
    assert "second:" in numbered[1].text
    assert _num_id_of(numbered[0]) == _num_id_of(numbered[1])


def test_block_selector_widget_sole_content_of_li_shares_one_numbered_paragraph(starter_template):
    """Regression (found visually in real content, question 235: a
    dropdown-rendered pl-multiple-choice, block display since is_dropdown
    disables the inline auto-detect signal): a block-display selector widget
    that's a <li>'s sole content must render ALL its options into the SAME
    numbered paragraph (joined by soft breaks), not one separate,
    unindented, unnumbered paragraph per option."""
    from pl2docx.html_parser import parse_instance_question_html

    html = """
    <div class="question-block">
      <div class="card-header"><h1>List of dropdown choices</h1></div>
      <div class="card-body question-body">
        <ol>
          <li>
            <span class="pl-multiple-choice-dropdown">
              <select name="q1">
                <option value="a" data-content="(a) chemical change">chemical change</option>
                <option value="b" data-content="(b) physical change">physical change</option>
              </select>
            </span>
          </li>
        </ol>
      </div>
    </div>
    """
    question = parse_instance_question_html(html)
    assert question.widgets[0].is_dropdown is True

    tpl = _tpl(starter_template)
    list_formats = create_list_formats(tpl)
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, question, 1, element_config, list_formats=list_formats)

    numbered = [p for p in ctx["question_contents"].paragraphs if _num_id_of(p) is not None]
    assert len(numbered) == 1  # one numbered paragraph, not one per option
    assert "chemical change" in numbered[0].text
    assert "physical change" in numbered[0].text
    assert "<w:br/>" in numbered[0]._p.xml  # options joined by a soft break, not real paragraphs


def test_each_minted_num_id_gets_start_override(starter_template):
    """Regression (found visually: answer-key list picked up numbering where
    the question's own list left off, e.g. 4/5/6 instead of restarting at
    1/2/3): every freshly-minted <w:num> must carry an explicit
    <w:lvlOverride><w:startOverride w:val="1"/></w:lvlOverride> - without it,
    Word doesn't reliably restart the count for <w:num> instances that share
    one <w:abstractNum>, even though each has its own numId."""
    from pl2docx.html_parser import parse_instance_question_html

    tpl = _tpl(starter_template)
    list_formats = create_list_formats(tpl)

    html = """
    <div class="question-block">
      <div class="card-header"><h1>List one</h1></div>
      <div class="card-body question-body">
        <ol><li>A</li><li>B</li></ol>
        <span class="input-group pl-integer-input">
          <input class="form-control pl-integer-input-input" name="answer" type="text">
        </span>
      </div>
    </div>
    """
    element_config = ElementConfig(preferences={}, behavior_class={})
    question = parse_instance_question_html(html)
    ctx1 = build_question_context(tpl, question, 1, element_config, list_formats=list_formats)
    ctx2 = build_question_context(tpl, question, 2, element_config, list_formats=list_formats)

    for ctx in (ctx1, ctx2):
        numbered = [p for p in ctx["question_contents"].paragraphs if _num_id_of(p) is not None]
        num_id = _num_id_of(numbered[0])
        numbering_xml = tpl.get_docx().part.numbering_part.element.xml
        num_match = re.search(
            rf'<w:num w:numId="{num_id}">.*?</w:num>', numbering_xml, re.DOTALL
        )
        assert num_match is not None
        assert "<w:startOverride" in num_match.group(0)
        assert 'w:val="1"' in num_match.group(0)


def test_block_selector_first_option_gets_leading_break_when_reusing_content(starter_template):
    """Regression (found visually): the FIRST option of a block-display
    selector widget must also get a soft break before it when reusing an
    already-populated numbered paragraph (e.g. image/arrow/colon text
    preceding it within the same <li>) - not just between options 2+."""
    from pl2docx.html_parser import parse_instance_question_html

    html = """
    <div class="question-block">
      <div class="card-header"><h1>Dropdown after other content</h1></div>
      <div class="card-body question-body">
        <ol>
          <li>Classify this:
            <span class="pl-multiple-choice-dropdown">
              <select name="q1">
                <option value="a" data-content="(a) chemical change">chemical change</option>
                <option value="b" data-content="(b) physical change">physical change</option>
              </select>
            </span>
          </li>
        </ol>
      </div>
    </div>
    """
    question = parse_instance_question_html(html)
    tpl = _tpl(starter_template)
    list_formats = create_list_formats(tpl)
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, question, 1, element_config, list_formats=list_formats)

    numbered = [p for p in ctx["question_contents"].paragraphs if _num_id_of(p) is not None]
    assert len(numbered) == 1
    lines = numbered[0].text.split("\n")
    assert lines[0].strip() == "Classify this:"
    assert lines[1].strip() == "○ chemical change"
    assert lines[2].strip() == "○ physical change"


def test_block_fill_in_gets_leading_break_when_reusing_content(starter_template):
    """Same general fix, for a fill-in (non-selector) block-display widget -
    the leading-break fix lives in the shared caller
    (_build_question_contents), not selector-specific code, so it must apply
    here too."""
    from pl2docx.html_parser import parse_instance_question_html

    html = """
    <div class="question-block">
      <div class="card-header"><h1>Fill-in after other content</h1></div>
      <div class="card-body question-body">
        <ol>
          <li>Enter a value:
            <span class="input-group pl-integer-input">
              <input class="form-control pl-integer-input-input" name="answer" type="text">
            </span>
          </li>
        </ol>
      </div>
    </div>
    """
    question = parse_instance_question_html(html)
    tpl = _tpl(starter_template)
    list_formats = create_list_formats(tpl)
    element_config = ElementConfig(
        preferences={"integer_input": FillInPreferences(display="block")}, behavior_class={}
    )
    ctx = build_question_context(tpl, question, 1, element_config, list_formats=list_formats)

    numbered = [p for p in ctx["question_contents"].paragraphs if _num_id_of(p) is not None]
    assert len(numbered) == 1
    lines = numbered[0].text.split("\n")
    assert lines[0].strip() == "Enter a value:"
    assert "Answer:" in lines[1]


def test_question_with_no_widgets_renders_successfully(starter_template):
    """Phase 5 subphase 3: a question with zero widgets (e.g. a diagram-only page,
    now supported by html_parser.py) must still produce a real, non-empty
    question_contents subdoc, an empty-but-present answer_element, and must not
    error - matching _build_question_contents' own `if not question.widgets:`
    fast path."""
    tpl = _tpl(starter_template)
    question = ParsedQuestion(
        title="Diagram-only question",
        prompt_segments=[[*plain("Here is a diagram:"), ImageRef(local_path="", alt="a diagram")]],
        widgets=[],
        answer_panel_text=None,
        points="1",
        points_numeric=1.0,
        qid="q/10",
    )
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, question, 1, element_config)

    assert "Here is a diagram" in _text(ctx["question_contents"])
    assert "a diagram" in _text(ctx["question_contents"])  # ImageRef alt-text fallback
    assert ctx["answer_element"] is not None
    assert _text(ctx["answer_element"]) == ""


def _svg_question(svg_markup='<svg viewBox="0 0 10 10"><rect width="10" height="10"/></svg>'):
    widget = Widget(kind="integer_input", name="answer")
    return ParsedQuestion(
        title="Q",
        prompt_segments=[
            [SvgRef(svg_markup=svg_markup, alt="a diagram")],
            plain(""),
        ],
        widgets=[widget],
        answer_panel_text=None,
        points="1",
        points_numeric=1.0,
        qid="q/8",
    )


def test_svg_ref_falls_back_to_placeholder_text_on_render_error(starter_template, monkeypatch):
    """Chromium unavailable (or a malformed SVG) must degrade to the alt
    text, not crash the whole document render - same philosophy as
    _render_math's LatexRenderError fallback."""
    import pl2docx.element_renderer as element_renderer_module

    def _boom(svg_markup, **kwargs):
        raise SvgRenderError("simulated failure")

    monkeypatch.setattr(element_renderer_module, "render_svg_png", _boom)

    tpl = _tpl(starter_template)
    question = _svg_question()
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, question, 1, element_config)

    assert _inline_shape_count(ctx["question_contents"]) == 0
    assert "a diagram" in _text(ctx["question_contents"])


def test_svg_image_ref_falls_back_when_file_missing(starter_template):
    """An <img src=*.svg> pointing at a file that isn't actually on disk must
    fall back gracefully, matching the existing missing-raster-image case."""
    tpl = _tpl(starter_template)
    widget = Widget(kind="integer_input", name="answer")
    question = ParsedQuestion(
        title="Q",
        prompt_segments=[
            [ImageRef(local_path="files/does_not_exist.svg", alt="missing diagram")],
            plain(""),
        ],
        widgets=[widget],
        answer_panel_text=None,
        points="1",
        points_numeric=1.0,
        qid="q/9",
    )
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, question, 1, element_config, image_base_dir=FIXTURES_DIR)

    assert _inline_shape_count(ctx["question_contents"]) == 0
    assert "missing diagram" in _text(ctx["question_contents"])


@pytest.mark.skipif(
    not chromium_available(), reason="requires Playwright's Chromium browser to be installed"
)
def test_svg_ref_embeds_as_real_picture(starter_template):
    """Phase 5 subphase 1: an inline SvgRef must become a real embedded (headless-
    browser-rasterized) picture, not the alt-text fallback, when Chromium is available."""
    tpl = _tpl(starter_template)
    question = _svg_question(
        '<svg viewBox="0 0 100 60" width="100" height="60" '
        'xmlns="http://www.w3.org/2000/svg"><line x1="10" y1="10" x2="90" y2="50" '
        'stroke="black" stroke-width="2"/></svg>'
    )
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, question, 1, element_config)

    assert _inline_shape_count(ctx["question_contents"]) == 1
    assert "a diagram" not in _text(ctx["question_contents"])


@pytest.mark.skipif(
    not chromium_available(), reason="requires Playwright's Chromium browser to be installed"
)
def test_svg_referenced_via_img_src_embeds_as_real_picture(starter_template):
    """Phase 5 subphase 1: an <img src=*.svg> (e.g. via <pl-figure>, confirmed to
    always emit a bare <img> regardless of file type) must rasterize and embed the
    same way a raw inline <svg> block does."""
    from pl2docx.html_parser import parse_instance_question_html

    html = (FIXTURES_DIR / "svg_img_blank.html").read_text(encoding="utf-8")
    question = parse_instance_question_html(html)

    tpl = _tpl(starter_template)
    element_config = ElementConfig(preferences={}, behavior_class={})
    ctx = build_question_context(tpl, question, 1, element_config, image_base_dir=FIXTURES_DIR)

    assert _inline_shape_count(ctx["question_contents"]) == 1
    assert "[image]" not in _text(ctx["question_contents"])
