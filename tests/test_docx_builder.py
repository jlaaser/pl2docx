import pytest
from docx import Document

from pl2docx.docx_builder import TemplateUnreadableError, render_document
from pl2docx.html_parser import ParsedQuestion, Widget, plain


def _mc_question(blank: bool) -> ParsedQuestion:
    widget = Widget(
        kind="multiple_choice",
        name="statement",
        options=[plain("Alpha"), plain("Beta"), plain("Gamma")],
        correct_option_indices=[] if blank else [1],
        is_inline=False,
    )
    return ParsedQuestion(
        title="MC Question",
        prompt_segments=[plain("Pick one."), plain("")],
        widgets=[widget],
        answer_panel_text=None if blank else plain("(B) Beta"),
        points="2",
        points_numeric=2.0,
        qid="course/questions/mc-question",
    )


def _integer_question(blank: bool) -> ParsedQuestion:
    widget = Widget(kind="integer_input", name="answer")
    return ParsedQuestion(
        title="Integer Question",
        prompt_segments=[plain("Enter a number."), plain("")],
        widgets=[widget],
        answer_panel_text=None if blank else plain("42"),
        points="1",
        points_numeric=1.0,
        qid="course/questions/integer-question",
    )


def _zones(blank: bool):
    return [
        {
            "title": "Zone One",
            "questions": [(_mc_question(blank), 1), (_integer_question(blank), 2)],
        },
        {
            "title": None,
            "questions": [],
        },
    ]


def _all_text(path):
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    # doc.paragraphs doesn't include paragraphs nested inside tables (e.g. the
    # bordered "SOLUTION:" box the starter template wraps answer_contents in).
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


def test_render_document_blank(starter_template, tmp_path):
    output_path = tmp_path / "blank.docx"
    zones = _zones(blank=True)
    # drop the empty second zone - not meaningful for a blank/key comparison test
    render_document(starter_template, [zones[0]], is_answer_key=False, output_path=output_path)

    assert output_path.exists()
    text = _all_text(output_path)
    assert "Zone One" in text
    assert "1. MC Question (2 points)" in text
    assert "Pick one." in text
    # multiple_choice with no inline signal auto-detects to block display, one option per line
    assert "◯ Alpha" in text
    assert "◯ Beta" in text
    assert "2. Integer Question (1 point)" in text
    assert "course/questions/mc-question" not in text  # qid only shown in the key
    assert "42" not in text


def test_render_document_threads_instance_id_into_template(starter_template, tmp_path):
    output_path = tmp_path / "blank.docx"
    zones = _zones(blank=True)
    render_document(
        starter_template, [zones[0]], is_answer_key=False, output_path=output_path,
        instance_id="Version A",
    )

    assert "Version A" in _all_text(output_path)


def test_render_document_key(starter_template, tmp_path):
    output_path = tmp_path / "key.docx"
    zones = _zones(blank=False)
    render_document(starter_template, [zones[0]], is_answer_key=True, output_path=output_path)

    text = _all_text(output_path)
    assert "[course/questions/mc-question]" in text
    assert "[course/questions/integer-question]" in text
    assert "42" in text

    doc = Document(str(output_path))
    beta_paragraph = next(p for p in doc.paragraphs if "Beta" in p.text)
    assert any(run.bold for run in beta_paragraph.runs if "Beta" in run.text)


def test_render_document_zone_without_title_has_no_heading(starter_template, tmp_path):
    output_path = tmp_path / "untitled_zone.docx"
    zone = {"title": None, "questions": [(_integer_question(blank=True), 1)]}
    render_document(starter_template, [zone], is_answer_key=False, output_path=output_path)

    doc = Document(str(output_path))
    heading_paragraphs = [p for p in doc.paragraphs if p.style.name == "pl2docx Zone Heading"]
    assert heading_paragraphs == []


def test_render_document_no_stray_blank_paragraphs_at_boundaries(starter_template, tmp_path):
    """Regression: the starter template's control-flow tags (for/if/else/endfor)
    previously left 2-5 stray blank paragraphs at every zone/question boundary,
    since a plain (un-prefixed) {% %} control tag's own paragraph survives
    rendering as a real empty <w:p> unless specially handled - confirmed fixed by
    switching every one of them to docxtpl's paragraph-consuming {%p %} syntax
    (see starter_template.py's module docstring for the full mechanism/why).
    Also covers a related but separate bug found while writing this test: a
    block-display widget with nothing after it left its own trailing empty
    paragraph inside question_contents itself (fixed in element_renderer.py's
    _trim_trailing_empty_paragraph, not a template issue).

    This fixture exercises every boundary type the user reported: before the
    zone loop, zone heading -> first question, question content -> answer key,
    and end of one zone's last question -> the next zone's heading. None of
    this fixture's real content is legitimately blank, so no paragraph in the
    rendered output should be empty."""
    output_path = tmp_path / "key.docx"
    zones = [
        {
            "title": "Zone One",
            "questions": [(_mc_question(blank=False), 1)],
        },
        {
            "title": "Zone Two",
            "questions": [(_integer_question(blank=False), 2)],
        },
    ]
    render_document(starter_template, zones, is_answer_key=True, output_path=output_path)

    doc = Document(str(output_path))
    texts = [p.text for p in doc.paragraphs]
    assert all(t.strip() for t in texts), f"found blank paragraph(s): {texts}"

    # Sanity-check the boundaries are actually adjacent to real content, not
    # just "no blanks anywhere" by coincidence of a lucky paragraph count.
    zone_one_idx = texts.index("Zone One")
    zone_two_idx = texts.index("Zone Two")
    assert texts[zone_one_idx - 1].startswith("Instructions:")
    assert texts[zone_one_idx + 1].startswith("1. MC Question")
    assert texts[zone_two_idx + 1].startswith("2. Integer Question")


def test_render_document_unreadable_template_raises_clear_error(tmp_path):
    """Regression: a template open in Word (or mid-cloud-sync) fails with a cryptic
    docx.opc.exceptions.PackageNotFoundError deep inside docxtpl's subdoc creation.
    Should surface as a clear, actionable pl2docx error instead."""
    not_a_docx = tmp_path / "not_a_docx.docx"
    not_a_docx.write_text("this is not a docx file", encoding="utf-8")

    with pytest.raises(TemplateUnreadableError):
        render_document(not_a_docx, [], is_answer_key=False, output_path=tmp_path / "out.docx")
