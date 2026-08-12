from docx import Document
from docxtpl import DocxTemplate

from pl2docx.starter_template import build_starter_template


def test_build_starter_template_defines_expected_styles(tmp_path):
    path = tmp_path / "template.docx"
    build_starter_template(path)

    assert path.exists()
    doc = Document(str(path))
    style_names = {s.name for s in doc.styles}
    assert "pl2docx Zone Heading" in style_names
    assert "pl2docx Question Title" in style_names
    assert "pl2docx QID Reference" in style_names


def test_build_starter_template_renders_with_empty_zones(tmp_path):
    """The generated template must itself be valid Jinja - no stray {{/{% look-alikes
    in the instructional boilerplate text (regression: an earlier version described the
    tag syntax in plain English inside the doc, which docxtpl then tried to parse as
    real Jinja and failed to compile)."""
    path = tmp_path / "template.docx"
    build_starter_template(path)

    tpl = DocxTemplate(str(path))
    tpl.render({"zones": [], "is_answer_key": False})
    output_path = tmp_path / "rendered.docx"
    tpl.save(str(output_path))
    assert output_path.exists()


def test_build_starter_template_solution_box_has_cell_border(tmp_path):
    """The "SOLUTION:" box must be a bordered table cell, not a bordered paragraph.

    Regression: a paragraph border on a {{p ... }} line is discarded, since
    docxtpl replaces that whole paragraph (border and all) with the substituted
    subdoc content. A table cell's border isn't affected, since it lives on the
    cell, not on the paragraphs inside it.
    """
    path = tmp_path / "template.docx"
    build_starter_template(path)

    doc = Document(str(path))
    assert len(doc.tables) == 1
    cell = doc.tables[0].cell(0, 0)
    assert "tcBorders" in cell._tc.xml
    cell_text = "\n".join(p.text for p in cell.paragraphs)
    assert "SOLUTION:" in cell_text
    assert "{{p question.answer_contents }}" in cell_text
